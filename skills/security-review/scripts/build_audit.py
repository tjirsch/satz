#!/usr/bin/env python3
"""Build the two Security-Toolset deliverables from Prowler OCSF + the CIS 5.0 catalog + satz evidence.

    uv run --with openpyxl --with python-docx --with pyyaml python3 build_audit.py \
        --scope scope.yaml --ocsf <merged.ocsf.json> [more...] \
        [--satz-report satz-report-cis-gcp-5.0.json] [--satz-require satz-require-cis-gcp-5.0.json] \
        [--evidence-dir <dir>] [--catalog cis_5.0_gcp.json] [--measures ../assets/measures.yaml] --out-dir <dir>

Outputs: CIS-5.0-Checkliste_<kunde>_<date>.xlsx and Remediation-Plan_<kunde>_CIS-5.0_<date>.docx.
Every one of the benchmark's controls gets a row and a status from the taxonomy
(FAIL / LÜCKE / MANUAL / PASS / FALSE POSITIVE / N/A (KEINE RESSOURCEN)); every FAIL/LÜCKE control
gets a measure from the catalogue; MANUAL controls get a review measure; N/A controls land in the
Wiedervorlage table. `scope.workloads` (the projects beside the estate) gets a "Workloads" section
in the plan and a "Workloads" sheet in the workbook, with each project's consumer check read from
the consumer-<name>.txt / checkov-<name>.json that satz_evidence.sh wrote beside the satz report
(--evidence-dir, default: the folder of --satz-report). Every number is computed from the inputs.
Nothing here touches the cloud, and no folder path reaches a document — a workload is named by its name.
"""

import argparse
import collections
import glob
import json
import re
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Cm

HERE = Path(__file__).resolve().parent
FAIL, LUECKE, MANUAL, PASS, FP, NA = (
    "FAIL",
    "LÜCKE",
    "MANUAL",
    "PASS",
    "FALSE POSITIVE",
    "N/A (KEINE RESSOURCEN)",
)
STATUS_ORDER = [FAIL, LUECKE, MANUAL, PASS, FP, NA]
FILLS = {
    MANUAL: ("FFE699", True, "7F6000"),
    PASS: ("C6EFCE", False, "006100"),
    NA: ("D0CECE", False, "3F3F3F"),
    FAIL: ("F8CBAD", True, "9C0006"),
    FP: ("BDD7EE", True, "1F3864"),
    LUECKE: ("FBC7C7", True, "9C0006"),
}
LESEHINWEISE = {
    FAIL: "Prowler hat einen Verstoß an vorhandenen Ressourcen festgestellt. Remediation im Plan (Spalte Maßnahme).",
    LUECKE: "Kein Scan-Ergebnis, weil die Grundlage fehlt — nicht weil die Kontrolle erfüllt wäre. Gilt als nicht erfüllt.",
    MANUAL: "Prowler prüft die Kontrolle nicht automatisiert. Stichprobe durch den CISO ist PFLICHT (Prüfweg in Spalte J).",
    PASS: "Prowler hat keinen Verstoß an vorhandenen Ressourcen gefunden. Stichprobe empfohlen bei Level-1-Kontrollen.",
    FP: "Prowler meldet FAIL, aber die Kontrolle ist nachweislich erfüllt (Begründung in Spalte Befund). Nachweis im Plan.",
    NA: "Der Ressourcentyp existiert in der Organisation nicht (kein Cloud SQL, kein BigQuery, …). Wiedervorlage beim Rollout.",
}
COLS = [
    "CIS #",
    "Kontrolle",
    "Sektion",
    "Level",
    "Prüfart",
    "Status (Prowler)",
    "Befund",
    "Betroffene Ressourcen",
    "Stichprobe CISO",
    "Prüfweg (Konsole / gcloud)",
    "Bewertung / Kommentar",
    "Maßnahme",
]
WIDTHS = [8, 46, 24, 8, 10, 20, 60, 36, 12, 70, 44, 12]
HEADER_FILL, SECTION_FILL, REVIEW_FILL = "1F3864", "D9E2F3", "FFF2CC"
PHASES = [
    ("sofort", "Sofort"),
    ("kurzfristig", "Kurzfristig"),
    ("nach_freigabe", "Nach Freigabe"),
    ("laufend", "Laufend"),
    ("review", "CISO-Review"),
]
RE_MAX = 12  # resources listed per control before "… +n weitere"
KIND_DE = {
    "satz": "satz-Projekt (eigenes Estate, liest das Interface)",
    "hcl": "HCL-Projekt (Terraform, liest module.satz)",
}
SKIPPED = 4  # satz_evidence.sh's exit code for a step it did not run, reason on the line before


class Safe(dict):
    def __missing__(self, k):
        return "{" + k + "}"


def fmt(text, ctx):
    return (text or "").format_map(Safe(ctx))


# --------------------------------------------------------------------------- inputs
def find_catalog():
    for pat in (
        "~/.local/share/uv/tools/prowler/lib/python3*/site-packages/prowler/compliance/gcp/cis_5.0_gcp.json",
        "/opt/homebrew/lib/python3*/site-packages/prowler/compliance/gcp/cis_5.0_gcp.json",
    ):
        hits = glob.glob(str(Path(pat).expanduser()))
        if hits:
            return hits[0]
    try:
        import prowler

        p = Path(prowler.__file__).parent / "compliance" / "gcp" / "cis_5.0_gcp.json"
        if p.exists():
            return str(p)
    except ImportError:
        pass
    sys.exit(
        "CIS 5.0 catalog not found — pass --catalog <prowler>/compliance/gcp/cis_5.0_gcp.json"
    )


def load_ocsf(paths):
    seen, out = set(), []
    for p in paths:
        for f in json.loads(Path(p).expanduser().read_text() or "[]"):
            uid = (
                f.get("finding_info", {}).get("uid")
                or json.dumps(f, sort_keys=True)[:200]
            )
            if uid not in seen:
                seen.add(uid)
                out.append(f)
    return out


def proj_of(f):
    return (f.get("cloud") or {}).get("account", {}).get("uid", "") or ""


def check_of(f):
    return (f.get("metadata") or {}).get("event_code", "") or ""


def cis_ids(f, version="5.0"):
    return (f.get("unmapped") or {}).get("compliance", {}).get(
        f"CIS-{version}", []
    ) or []


def resource_of(f):
    r = (f.get("resources") or [{}])[0]
    return r.get("name") or r.get("uid") or ""


def detail(f):
    return (
        f.get("status_detail") or f.get("finding_info", {}).get("title") or ""
    ).strip()


def fill_placeholders(text, scope, project):
    rep = {
        "ORGANIZATION_ID": scope["org_id"],
        "ORG_ID": scope["org_id"],
        "PROJECT_ID": project,
        "PROJECT_NAME": project,
        "DOMAIN": scope["domain"],
        "YOUR_DOMAIN": scope["domain"],
    }
    for k, v in rep.items():
        text = re.sub(r"[<\[{]?\b" + k + r"\b[>\]}]?", str(v), text)
    return text


# --------------------------------------------------------------------------- workloads
def strip_paths(line):
    """A local path down to its last component — a document names no folder (PII) — and the
    finding layout's padded columns down to one space."""
    line = re.sub(r"(?:/Users/|/home/|~/|\$HOME/)[^\s:]*/", "", line)
    return re.sub(r"\s{2,}", " ", line)


def read_consumer_file(path):
    """consumer-<name>.txt → [{title, lines, exit}]: one block per `== step`, closed by `exit: <code>`."""
    blocks = []
    for line in path.read_text(errors="replace").splitlines():
        if line.startswith("== "):
            blocks.append(dict(title=line[3:].strip(), lines=[], exit=None))
        elif line.startswith("exit: ") and blocks:
            blocks[-1]["exit"] = int(line[6:].split()[0])
        elif blocks:
            blocks[-1]["lines"].append(line)
    return blocks


def checkov_counts(path):
    reports = json.loads(path.read_text() or "[]")
    reports = reports if isinstance(reports, list) else [reports]
    return {
        k: sum((r.get("summary") or {}).get(k, 0) for r in reports)
        for k in ("passed", "failed")
    }


def load_workloads(scope, evidence_dir):
    """scope.workloads joined with the evidence satz_evidence.sh wrote:
    [{name, kind, kind_de, project_id, interface, check, check_ok (True/False/None), checkov}]."""
    out = []
    for w in scope.get("workloads") or []:
        name = w["name"]
        row = dict(
            name=name,
            kind=w.get("kind") or "",
            kind_de=KIND_DE.get(w.get("kind"), w.get("kind") or "—"),
            project_id=w.get("project_id") or "",
            interface=w.get("interface") or "",
        )
        f = evidence_dir / f"consumer-{name}.txt" if evidence_dir else None
        if f is None or not f.exists():
            row.update(
                check="kein Nachweis — satz_evidence.sh ist ohne --workloads gelaufen",
                check_ok=None,
                checkov="nicht gelaufen",
            )
            out.append(row)
            continue
        blocks = read_consumer_file(f)
        checks = [b for b in blocks if not b["title"].startswith("checkov")]
        ck = [b for b in blocks if b["title"].startswith("checkov")]
        problems = []
        for b in checks:
            if b["exit"] == 0:
                continue
            step = b["title"].split(" ")[0]
            notable = [
                strip_paths(ln.strip())[:200]
                for ln in b["lines"]
                if ln.strip()
                and (
                    ln.lower().startswith("error") or step in ln or b["exit"] == SKIPPED
                )
            ]
            problems.append(
                f"{step}: " + ("; ".join(notable[:4]) or f"exit {b['exit']}")
            )
        row["check_ok"] = bool(checks) and not problems
        if row["check_ok"]:
            row["check"] = (
                "bestanden — "
                + ", ".join(b["title"].split(" ")[0] for b in checks)
                + " ohne Befund"
            )
        else:
            row["check"] = " | ".join(problems) or "kein Prüfschritt protokolliert"
        cj = evidence_dir / f"checkov-{name}.json"
        if ck and ck[0]["exit"] == SKIPPED:
            note = " ".join(ln.strip() for ln in ck[0]["lines"] if ln.strip())
            row["checkov"] = "nicht gelaufen — " + strip_paths(note)[:160]
        elif cj.exists():
            c = checkov_counts(cj)
            row["checkov"] = f"{c['failed']} Befunde, {c['passed']} bestanden"
        else:
            row["checkov"] = "nicht gelaufen"
        out.append(row)
    return out


# --------------------------------------------------------------------------- model
def build_model(scope, catalog, findings, satz_rows, require):
    fleet_re = re.compile(scope["fleet_regex"]) if scope.get("fleet_regex") else None
    is_fleet = (lambda p: bool(fleet_re.match(p))) if fleet_re else (lambda p: False)
    projects = sorted({proj_of(f) for f in findings if proj_of(f)})
    own = [p for p in projects if not is_fleet(p)]
    fleet = [p for p in projects if is_fleet(p)]
    fleet_label = re.sub(r"^\s*\d+\s+", "", scope.get("fleet_label") or "") or (
        "generierte Projekte" if fleet else ""
    )
    for cid in scope.get("luecke") or {}:
        req = next((r for r in catalog["Requirements"] if r["Id"] == cid), None)
        if (
            req is None
            or req["Attributes"][0]["AssessmentStatus"] == "Manual"
            or not req["Checks"]
        ):
            print(
                f"WARNING: luecke[{cid}] ignored — control is MANUAL (or unknown); MANUAL outranks LÜCKE in the taxonomy",
                file=sys.stderr,
            )
    overrides = scope.get("check_overrides") or {}
    by_control = collections.defaultdict(list)
    for f in findings:
        ids = list(cis_ids(f)) + list(overrides.get(check_of(f), []))
        for cid in dict.fromkeys(ids):
            by_control[cid].append(f)
    satz_by = {r["control"]: r for r in satz_rows}
    req_by = {c["id"]: c for c in (require or {}).get("controls", [])}
    compute_scanned = any(check_of(f).startswith("compute_") for f in findings)
    luecke = dict(scope.get("luecke") or {})
    # automatic LÜCKE: the sink-lock rows without an org sink, and every compute-backed control if compute never ran
    for cid, txt in (
        (
            "2.3",
            "kein Org-Log-Sink vorhanden — Prowler betrachtet nur Sinks nach storage.googleapis.com",
        ),
        (
            "2.4",
            "kein Org-Log-Sink vorhanden — Aufbewahrungsrichtlinie kann nicht geprüft werden",
        ),
    ):
        luecke.setdefault(cid, txt)
    rows = []
    for req in catalog["Requirements"]:
        cid, a = req["Id"], req["Attributes"][0]
        fs = by_control.get(cid, [])
        fails = [f for f in fs if f.get("status_code") == "FAIL"]
        passes = [f for f in fs if f.get("status_code") == "PASS"]
        own_fails = [f for f in fails if not is_fleet(proj_of(f))]
        fleet_fails = [f for f in fails if is_fleet(proj_of(f))]
        compute_backed = any(c.startswith("compute_") for c in req["Checks"])
        if cid in (scope.get("false_positives") or {}):
            status = FP
        elif fails:
            status = FAIL
        elif passes:
            status = PASS
        elif a["AssessmentStatus"] == "Manual" or not req["Checks"]:
            status = MANUAL
        elif cid in luecke and cid in (scope.get("luecke") or {}):
            status = LUECKE
        elif compute_backed and not compute_scanned:
            status = LUECKE
            luecke[cid] = (
                "Compute nicht gescannt (eigene Projekte nicht übergeben) — kein Ergebnis"
            )
        elif cid in ("2.3", "2.4") and (
            satz_by.get(cid, {}).get("status") in (None, "unmet")
        ):
            status = LUECKE
        else:
            status = NA
        # Befund
        parts = []
        if own_fails:
            seen = set()
            for f in own_fails:
                d = detail(f)
                if d not in seen:
                    seen.add(d)
                    parts.append(d)
            if len(parts) > 8:
                parts = parts[:8] + [f"… +{len(parts) - 8} weitere Befunde"]
        if fleet_fails:
            nfp = len({proj_of(f) for f in fleet_fails})
            parts.append(
                f"{nfp} {fleet_label}: {detail(fleet_fails[0]).replace(proj_of(fleet_fails[0]), '<Projekt>')}"
            )
        if scope.get("measures") and not isinstance(scope["measures"], dict):
            sys.exit(
                "scope.measures must be a mapping {skip, override, extra} — an old-format scope (measures: M1: [...]) needs converting"
            )
        if status == PASS:
            parts = [detail(f) for f in passes[:3]] + (["…"] if len(passes) > 3 else [])
        elif status == LUECKE:
            parts = [luecke.get(cid, "Grundlage fehlt")]
        elif status == FP:
            parts = [scope["false_positives"][cid]]
        elif status == MANUAL:
            parts = [
                "Manual check — Prowler prüft diese Kontrolle nicht; Stichprobe nach Prüfweg"
            ]
        elif status == NA:
            parts = ["kein Ergebnis — Ressourcentyp nicht vorhanden"]
        res_own = sorted({f"{resource_of(f)} ({proj_of(f)})" for f in own_fails})
        res = res_own[:RE_MAX] + (
            [f"… +{len(res_own) - RE_MAX} weitere"] if len(res_own) > RE_MAX else []
        )
        if fleet_fails:
            res.append(f"{len({proj_of(f) for f in fleet_fails})} Flotten-Projekte")
        proj_for = proj_of(own_fails[0]) if own_fails else scope["infra_project"]
        pruefweg = fill_placeholders(a.get("AuditProcedure") or "", scope, proj_for)
        # Bewertung seed from satz (declared side) + require (goal view)
        sr, rq, bew = satz_by.get(cid), req_by.get(cid), ""
        if sr:
            st = sr.get("status", "")
            wit = ", ".join(
                (w.get("address") or w.get("declared") or "")
                if isinstance(w, dict)
                else str(w)
                for w in (sr.get("witnesses") or [])[:2]
            )
            if st == "verified":
                bew = "[satz] deklariert und live verifiziert" + (
                    f" — {wit}" if wit else ""
                )
            elif st == "DRIFTED":
                bew = (
                    "[satz] deklariert, live abweichend — Apply des Estates ausstehend"
                )
            elif st.startswith("partial"):
                duties = sr.get("duties") or ""
                bew = (
                    "[satz] deklariert (trägt bei), organisatorischer Rest offen"
                    if "contributes" in st or duties in ("", "–")
                    else f"[satz] deklariert, offene Pflicht: {duties}"
                )
            elif st == "unmet":
                bew = "[satz] nicht deklariert"
            elif st == "organizational":
                bew = "[satz] organisatorische Kontrolle (Kunde)"
            if sr.get("checkov") not in (None, "", "–"):
                bew += f" · Checkov: {sr['checkov']}"
        if rq and rq.get("providers") and rq.get("verdict") in ("unmet", "partial"):
            bew += (
                (" — " if bew else "[satz] ")
                + "Pack stellt es bereit: "
                + ", ".join(rq["providers"])
            )
        rows.append(
            dict(
                cid=cid,
                title=req["Description"],
                section=a["Section"],
                level=a["Profile"],
                pruefart=a["AssessmentStatus"],
                status=status,
                befund=" | ".join(parts),
                resources="; ".join(res),
                stichprobe="PFLICHT"
                if status == MANUAL
                else ("empfohlen" if status in (FAIL, LUECKE) else ""),
                pruefweg=pruefweg,
                bewertung=bew,
                massnahme="",
                checks=req["Checks"],
                own_fails=own_fails,
                fleet_fails=fleet_fails,
                passes=passes,
                remediation=a.get("RemediationProcedure") or "",
                satz=sr,
                require=rq,
                res_own=res_own,
            )
        )
    counts = collections.Counter(r["status"] for r in rows)
    return dict(
        rows=rows,
        counts=counts,
        own=own,
        fleet=fleet,
        fleet_label=fleet_label,
        findings=findings,
        by_control=by_control,
        compute_scanned=compute_scanned,
    )


# --------------------------------------------------------------------------- measures
def select_measures(scope, model, catalogue):
    rows = {r["cid"]: r for r in model["rows"]}
    cfg = scope.get("measures") or {}
    skip = set(cfg.get("skip") or [])
    chosen = []
    for m in catalogue + list(cfg.get("extra") or []):
        m = dict(m)
        m.update((cfg.get("override") or {}).get(m["id"], {}))
        if m["id"] in skip:
            continue
        if m.get("fleet_only"):
            if not model["fleet"]:
                continue
            m["controls"] = sorted(
                {r["cid"] for r in model["rows"] if r["fleet_fails"]}, key=ver_key
            )
        ctrls = [rows[c] for c in m.get("controls", []) if c in rows]
        active = [r for r in ctrls if r["status"] in (FAIL, LUECKE, FP)]
        drifted = [r for r in ctrls if (r.get("satz") or {}).get("status") == "DRIFTED"]
        manual = [r for r in ctrls if r["status"] == MANUAL]
        if (
            m.get("always")
            or active
            or (m["kind"] == "satz-apply" and drifted)
            or (m["kind"] == "review" and manual)
        ):
            m["_active"] = active
            m["_drifted"] = drifted
            m["_manual"] = manual
            m["_rows"] = ctrls
            chosen.append(m)
    order = {p: i for i, (p, _) in enumerate(PHASES)}
    chosen.sort(
        key=lambda m: order.get(m["phase"], 99)
    )  # stable: catalogue order within a phase
    for i, m in enumerate(chosen, 1):
        m["num"] = f"M{i}"
        for r in m["_rows"]:
            if r["status"] in (FAIL, LUECKE, FP, MANUAL) or r in m["_drifted"]:
                r["massnahme"] = (r["massnahme"] + ", " if r["massnahme"] else "") + m[
                    "num"
                ]
    # every FAIL/LÜCKE control must have a measure — otherwise synthesise one from the CIS remediation text
    orphans = [
        r for r in model["rows"] if r["status"] in (FAIL, LUECKE) and not r["massnahme"]
    ]
    for r in orphans:
        n = f"M{len(chosen) + 1}"
        m = dict(
            id=f"auto-{r['cid']}",
            num=n,
            title=r["title"],
            phase="kurzfristig",
            kind="gcloud",
            controls=[r["cid"]],
            aufwand="siehe CIS-Remediation",
            risiko="siehe CIS-Remediation",
            why="Kontrolle ohne Katalog-Maßnahme — Remediation aus dem CIS-Benchmark.",
            gcloud=r["remediation"],
            _active=[r],
            _drifted=[],
            _manual=[],
            _rows=[r],
        )
        chosen.append(m)
        r["massnahme"] = n
    by_m = {c: m for m in catalogue for c in m.get("controls", [])}
    for r in model["rows"]:
        if r["status"] == NA and not r["massnahme"]:
            r["massnahme"] = "Wiedervorlage"
        elif r["status"] == MANUAL and not r["massnahme"]:
            r["massnahme"] = "Review"
            r["review_ref"] = by_m[r["cid"]]["title"] if r["cid"] in by_m else ""
    return chosen


def ver_key(cid):
    return [int(x) for x in cid.split(".")]


def measure_ctx(scope, model, m):
    own_p = sorted({proj_of(f) for r in m["_rows"] for f in r["own_fails"]})
    resources = sorted({x for r in m["_rows"] for x in r["res_own"]})
    all_p = sorted(
        {proj_of(f) for r in m["_rows"] for f in r["own_fails"] + r["fleet_fails"]}
    )
    ctx = dict(
        org_id=scope["org_id"],
        domain=scope["domain"],
        kunde=scope["kunde_lang"],
        shortcode=scope["kunde"],
        estate_repo=Path(scope["estate_repo"]).name,
        estate_file=scope.get("estate_file", "<estate>.satz"),
        iac_sa=scope.get("iac_sa", ""),
        infra_project=scope["infra_project"],
        n_own=len(model["own"]),
        n_fleet=len(model["fleet"]),
        n_projects=len(all_p) or len(model["own"]) + len(model["fleet"]),
        fleet_label=model["fleet_label"],
        fleet_regex=scope.get("fleet_regex", ""),
        projects=" ".join(own_p) or "<projekt>",
        resources=", ".join(resources[:RE_MAX]) or "—",
        log_project=f"{scope['kunde']}-log-infra-001",
        default_region=scope.get("default_region", "europe-west3"),
        retention=scope.get("retention_days", 400),
        controls_line=", ".join(r["cid"] for r in m["_active"])
        or ", ".join(m.get("controls", [])),
    )
    return ctx


# --------------------------------------------------------------------------- xlsx
def build_xlsx(scope, model, measures, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Übersicht"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 58

    def put(r, a, b=None, bold=False, size=11, fill=None):
        ws.cell(row=r, column=1, value=a).font = Font(bold=bold, size=size)
        if b is not None:
            c = ws.cell(row=r, column=2, value=b)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)

    put(1, f"{scope['framework']} — Statusmatrix", bold=True, size=14)
    put(3, "Kunde", scope["kunde_lang"], bold=True)
    put(4, "Organisation", f"{scope['domain']} (Org {scope['org_id']})", bold=True)
    put(5, "Scan-Datum", scope["scan_date_de"], bold=True)
    put(6, "Scanner", scope["scanner"], bold=True)
    scope_txt = f"{len(model['own'])} eigene Projekte ({', '.join(model['own'])})"
    if model["fleet"]:
        scope_txt += f" + {len(model['fleet'])} {model['fleet_label']} — aggregiert"
    put(7, "Gescannte Projekte", scope_txt, bold=True)
    if "workloads" in scope:
        wl = model["workloads"]
        put(
            8,
            "Workloads",
            (
                f"{len(wl)} neben dem Estate ("
                + ", ".join(
                    w["name"] + (f" · {w['project_id']}" if w["project_id"] else "")
                    for w in wl
                )
                + ") — Blatt »Workloads«"
            )
            if wl
            else "keine neben dem Estate",
            bold=True,
        )
    n = len(model["rows"])
    put(9, f"Statusverteilung ({n} Kontrollen)", bold=True, size=12)
    put(10, "Status", "Anzahl", bold=True)
    r = 11
    for st in STATUS_ORDER:
        put(r, st, f"=COUNTIF('CIS 5.0 Checkliste'!F:F,\"{st}\")", fill=FILLS[st][0])
        r += 1
    put(r, "Summe", f"=SUM(B11:B{r - 1})", bold=True)
    r += 2
    put(r, "Lesehinweise", bold=True, size=12)
    r += 1
    for st in STATUS_ORDER:
        put(r, st, LESEHINWEISE[st], fill=FILLS[st][0])
        r += 1
    if scope.get("scan_caveats"):
        r += 1
        put(r, "Werkzeuggrenzen dieses Scans", bold=True, size=12)
        r += 1
        for c in scope["scan_caveats"]:
            put(r, "•", c)
            r += 1
    r += 1
    put(r, "Prüfweg", bold=True, size=12)
    r += 1
    put(
        r,
        "Spalte »Prüfweg (Konsole / gcloud)« enthält die Audit-Prozedur aus dem CIS-Benchmark, mit den Werten dieser Umgebung (Org-ID, Projekt) befüllt — direkt nachprüfbar, ohne den Benchmark aufzuschlagen.",
    )
    r += 2
    put(r, "Auszufüllen im Review", bold=True, size=12)
    r += 1
    put(
        r,
        "Gelb hinterlegte Spalten der Checkliste: »Bewertung / Kommentar« (Vorbelegung »[satz] …« = was das Estate deklariert und live verifiziert) und »Maßnahme« (Mn = Abschnitt im Remediation-Plan; »Wiedervorlage« = Ressourcentyp nicht vorhanden).",
    )
    r += 2
    put(r, "Maßnahmen im Remediation-Plan", bold=True, size=12)
    r += 1
    for m in measures:
        put(r, m["num"], f"{m['title']} — {', '.join(m.get('controls', []))}")
        r += 1
    r += 1
    put(r, "Beispielzeile (Format)", bold=True, size=12)
    r += 1
    put(
        r,
        'Bewertung: "Bucket nur intern erreichbar, Risiko akzeptiert (Ticket SEC-123)" · Maßnahme: "M2 — bis 30.09."',
    )

    ws = wb.create_sheet("CIS 5.0 Checkliste")
    for i, (h, w) in enumerate(zip(COLS, WIDTHS), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    r, last = 2, None
    for row in model["rows"]:
        if row["section"] != last:
            c = ws.cell(row=r, column=1, value=row["section"])
            c.font = Font(bold=True, size=10)
            for i in range(1, 13):
                ws.cell(row=r, column=i).fill = PatternFill(
                    "solid", fgColor=SECTION_FILL
                )
            last = row["section"]
            r += 1
        vals = [
            row["cid"],
            row["title"],
            row["section"],
            row["level"],
            row["pruefart"],
            row["status"],
            row["befund"],
            row["resources"],
            row["stichprobe"],
            row["pruefweg"],
            row["bewertung"],
            row["massnahme"],
        ]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(size=10)
            c.alignment = Alignment(
                wrap_text=i in (2, 7, 8, 10, 11, 12), vertical="top"
            )
        st = ws.cell(row=r, column=6)
        f, b, col = FILLS[row["status"]]
        st.fill = PatternFill("solid", fgColor=f)
        st.font = Font(size=10, bold=b, color=col)
        st.alignment = Alignment(vertical="center")
        for i in (11, 12):
            ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=REVIEW_FILL)
        ws.row_dimensions[r].height = (
            108 if row["status"] in (FAIL, LUECKE, MANUAL) else 30
        )
        r += 1
    ws.auto_filter.ref = f"A1:L{r - 1}"

    ws = wb.create_sheet("Findings (FAIL)")
    fcols = [
        "CIS #",
        "Check-ID",
        "Severity",
        "Projekt",
        "Ressource",
        "Befund",
        "Remediation (Kurz)",
    ]
    for i, (h, w) in enumerate(zip(fcols, [8, 34, 11, 20, 26, 46, 68]), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    r = 2
    for row in model["rows"]:
        for f in row["own_fails"]:
            vals = [
                row["cid"],
                check_of(f),
                f.get("severity", ""),
                proj_of(f),
                resource_of(f),
                detail(f),
                (f.get("remediation") or {}).get("desc", "")[:600],
            ]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = Font(size=10)
                c.alignment = Alignment(wrap_text=i in (6, 7), vertical="top")
            r += 1
        if row["fleet_fails"]:
            f = row["fleet_fails"][0]
            nproj = len({proj_of(x) for x in row["fleet_fails"]})
            vals = [
                row["cid"],
                check_of(f),
                f.get("severity", ""),
                f"{nproj} Projekte (Flotte)",
                model["fleet_label"],
                detail(f).replace(proj_of(f), "<Projekt>"),
                (f.get("remediation") or {}).get("desc", "")[:600],
            ]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = Font(size=10, italic=True)
                c.alignment = Alignment(wrap_text=i in (6, 7), vertical="top")
            r += 1
    ws.auto_filter.ref = f"A1:G{max(r - 1, 2)}"

    if "workloads" in scope:
        ws = wb.create_sheet("Workloads")
        wcols = [
            "Workload",
            "Art",
            "Google-Projekt",
            "Interface",
            "Consumer-Check",
            "Checkov",
        ]
        for i, (h, w) in enumerate(zip(wcols, [18, 42, 28, 14, 72, 30]), 1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, size=10, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=HEADER_FILL)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        r = 2
        for w in model["workloads"]:
            vals = [
                w["name"],
                w["kind_de"],
                w["project_id"] or "—",
                w["interface"] or "—",
                w["check"],
                w["checkov"],
            ]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = Font(size=10)
                c.alignment = Alignment(wrap_text=i in (2, 5, 6), vertical="top")
            f, b, col = FILLS[
                PASS if w["check_ok"] else (NA if w["check_ok"] is None else FAIL)
            ]
            ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=f)
            ws.cell(row=r, column=5).font = Font(size=10, bold=b, color=col)
            r += 1
        if not model["workloads"]:
            ws.cell(
                row=2,
                column=1,
                value="keine Workloads neben dem Estate (scope.yaml: workloads: [])",
            )
    wb.save(out)


# --------------------------------------------------------------------------- docx helpers
def add_table(doc, rows, widths=None, header=False):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, v in enumerate(row):
            cells[ci].text = ""
            run = cells[ci].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
            if (header and ri == 0) or (not header and ci == 0):
                run.bold = True
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Cm(w)
    doc.add_paragraph()
    return t


def add_code(doc, text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    doc.add_paragraph()


def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def status_line(m):
    parts = []
    for r in m["_rows"]:
        if r["status"] == FAIL:
            n_own = len({proj_of(f) for f in r["own_fails"]})
            n_fl = len({proj_of(f) for f in r["fleet_fails"]})
            where = " · ".join(
                x
                for x in [
                    f"{n_own} eigene Projekte" if n_own else "",
                    f"{n_fl} Flotten-Projekte" if n_fl else "",
                ]
                if x
            )
            parts.append(f"{r['cid']} FAIL ({where})")
        elif r["status"] in (LUECKE, FP) or (
            r["status"] == MANUAL and m["kind"] == "review"
        ):
            parts.append(f"{r['cid']} {r['status']}")
        elif r in m["_drifted"]:
            parts.append(f"{r['cid']} DRIFTED (satz)")
    return "; ".join(parts) or "vorbeugend"


# --------------------------------------------------------------------------- docx
def build_docx(scope, model, measures, out):
    cnt, own, nfleet = model["counts"], model["own"], len(model["fleet"])
    S, ORG, DOM = scope, scope["org_id"], scope["domain"]
    n_own_f = sum(len(r["own_fails"]) for r in model["rows"])
    n_fl_f = sum(len(r["fleet_fails"]) for r in model["rows"])
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Remediation-Plan", level=0)
    para(doc, S["framework"], bold=True)
    scope_txt = (
        f"{len(own)} eigene Projekte"
        + (f" + {nfleet} {model['fleet_label']}" if nfleet else "")
        + f" — {len(own) + nfleet} Projekte unter der Organisation"
    )
    add_table(
        doc,
        [
            ["Kunde", S["kunde_lang"]],
            ["Organisation", f"{DOM} (Org {ORG})"],
            ["Scan-Datum", S["scan_date_de"]],
            ["Scanner", S["scanner"]],
            ["Scope", scope_txt],
            ["Status", S.get("estate_status", "")],
        ],
        widths=[3.5, 13.5],
    )
    para(
        doc,
        "Umsetzung: Die beschriebenen Maßnahmen werden ausschließlich durch den Kunden bzw. den Betrieb ausgeführt. "
        "Dieser Plan liefert geprüfte Kommandos und Konfigurationen, führt aber selbst keine Änderungen durch. "
        "Wo eine Maßnahme als satz-Pack vorliegt, ist das die bevorzugte Form: einmal deklariert, per Apply ausgerollt, "
        "durch report-compliance laufend nachgewiesen.",
        italic=True,
    )

    # 1 Ausgangslage
    doc.add_heading("1  Ausgangslage", level=1)
    txt = f"Der Scan umfasst {len(own) + nfleet} Projekte: {len(own)} eigene Projekte"
    if nfleet:
        txt += f" und {nfleet} {model['fleet_label']}, die in allen Kontrollen dasselbe Bild zeigen und deshalb als eine Gruppe geführt werden"
    txt += f". Bewertet wurden alle {len(model['rows'])} Kontrollen des Benchmarks gegen den vollständigen Katalog, nicht nur die, die der Scan gemeldet hat."
    para(doc, txt)
    verified = [
        r["cid"]
        for r in model["rows"]
        if (r.get("satz") or {}).get("status") == "verified"
    ]
    drifted = [
        r["cid"]
        for r in model["rows"]
        if (r.get("satz") or {}).get("status") == "DRIFTED"
    ]
    add_table(
        doc,
        [
            ["Status", "Anzahl", "Bewertung"],
            [
                FAIL,
                cnt[FAIL],
                f"Verstöße an vorhandenen Ressourcen — Maßnahmen M1–M{len(measures)}",
            ],
            [
                LUECKE,
                cnt[LUECKE],
                "Grundlage fehlt — nach der zugehörigen Maßnahme erstmals prüfbar",
            ],
            [
                MANUAL,
                cnt[MANUAL],
                "Klärung im CISO-Review (Pflicht-Stichprobe nach Prüfweg)",
            ],
            [
                PASS,
                cnt[PASS],
                "erfüllt"
                + (
                    f" — {len(verified)} davon durch deklarierte und live verifizierte Ressourcen des Estates"
                    if verified
                    else ""
                ),
            ],
            [
                FP,
                cnt[FP],
                "Prowler-FAIL, nachweislich erfüllt (Begründung in der Checkliste)",
            ],
            [
                NA,
                cnt[NA],
                "Ressourcentyp nicht vorhanden — Wiedervorlage beim Workload-Rollout (Abschnitt Wiedervorlage)",
            ],
        ],
        widths=[4.5, 2, 10.5],
        header=True,
    )
    para(doc, "Was die Zahlen bedeuten", bold=True)
    bullets = []
    log_rows = [
        r
        for r in model["rows"]
        if r["cid"].startswith("2.") and r["status"] in (FAIL, LUECKE)
    ]
    if len(log_rows) >= 6:
        bullets.append(
            f"Die Logging-Kontrollen (2.x) schlagen in {len(log_rows)} Zeilen an — Kennzeichen eines fehlenden Logging-Fundaments (Org-Sink, zentrales Alerting). "
            f"Eine Pack-Maßnahme schließt sie für alle {len(own) + nfleet} Projekte gleichzeitig; Prowler rechnet org-weites Alerting den Kindprojekten an."
        )
    if verified:
        bullets.append(
            (
                f"{len(verified)} Kontrollen ({', '.join(verified)}) sind"
                if len(verified) > 1
                else f"Kontrolle {verified[0]} ist"
            )
            + " durch das Estate deklariert und live verifiziert — präventiv für jedes neue Projekt."
        )
    if drifted:
        bullets.append(
            (
                f"{len(drifted)} deklarierte Kontrollen ({', '.join(drifted)}) weichen"
                if len(drifted) > 1
                else f"Die deklarierte Kontrolle {drifted[0]} weicht"
            )
            + " live ab — das Estate ist konvergiert, aber nicht angewendet (M1)."
        )
    bullets.append(
        f"{n_own_f} Einzelbefunde in eigenen Projekten"
        + (f" und {n_fl_f} in den generierten Projekten" if nfleet else "")
        + f" sind keine {n_own_f + n_fl_f} Probleme: hinter ihnen stehen {len(measures)} Maßnahmen."
    )
    add_bullets(doc, bullets)

    # 2 Reihenfolge
    doc.add_heading("2  Reihenfolge der Maßnahmen", level=1)
    lines = []
    for key, label in PHASES:
        ms = [m for m in measures if m["phase"] == key]
        if ms:
            lines.append(
                f"{label}: " + "; ".join(f"{m['num']} — {m['title']}" for m in ms)
            )
    add_bullets(doc, lines)
    para(
        doc,
        f"Alle satz-Kommandos laufen als IaC-Service-Account des Estates ({S.get('iac_sa', '')}); gcloud-Alternativen brauchen eine Identität mit "
        "Schreibrechten auf Organisationsebene — nicht den read-only Audit-Zugang.",
    )

    # measures
    sec = 3
    for m in measures:
        ctx = measure_ctx(scope, model, m)
        doc.add_heading(f"{sec}  {m['num']}  {m['title']}", level=1)
        sec += 1
        ctrl = (
            " · ".join(f"{r['cid']} ({r['level']})" for r in m["_rows"])
            or "— (ohne CIS-Kontrolle)"
        )
        add_table(
            doc,
            [
                ["CIS-Kontrolle", ctrl],
                ["Status", status_line(m)],
                ["Aufwand", fmt(m.get("aufwand", ""), ctx)],
                ["Risiko", fmt(m.get("risiko", ""), ctx)],
            ],
            widths=[3.5, 13.5],
        )
        bef = []
        for r in m["_rows"]:
            if (
                r["status"] in (FAIL, LUECKE, FP)
                or r in m["_drifted"]
                or (m["kind"] == "review" and r["status"] == MANUAL)
            ):
                line = f"{r['cid']} {r['title']}: " + (
                    r["bewertung"]
                    if r in m["_drifted"] and r["status"] not in (FAIL, LUECKE)
                    else r["befund"][:400]
                )
                if r["res_own"]:
                    line += f" — {', '.join(r['res_own'][:6])}" + (
                        f" … +{len(r['res_own']) - 6}" if len(r["res_own"]) > 6 else ""
                    )
                bef.append(line)
        if bef:
            para(doc, "Befund", bold=True)
            add_bullets(doc, bef)
        para(doc, fmt(m.get("why", ""), ctx))
        provs = sorted(
            {
                p
                for r in m["_rows"]
                for p in ((r.get("require") or {}).get("providers") or [])
            }
        )
        if provs:
            para(
                doc,
                "satz require: nicht im Estate deklariert — bereitgestellt durch Pack "
                + ", ".join(provs)
                + ".",
                italic=True,
            )
        if m.get("satz"):
            para(doc, "Umsetzung mit satz", bold=True)
            add_code(doc, fmt(m["satz"], ctx))
        if m.get("gcloud"):
            para(
                doc,
                "Alternative ohne satz (gcloud)"
                if m.get("satz")
                else "Umsetzung (gcloud)",
                bold=True,
            )
            add_code(doc, fmt(m["gcloud"], ctx))
        if m.get("terraform"):
            para(doc, "Terraform", bold=True)
            add_code(doc, fmt(m["terraform"], ctx))
        if not m.get("satz") and not m.get("gcloud"):
            for r in m["_rows"]:
                if r["remediation"]:
                    para(doc, f"CIS-Remediation {r['cid']}", bold=True)
                    add_code(
                        doc,
                        fill_placeholders(
                            r["remediation"], scope, scope["infra_project"]
                        ),
                    )
        if m.get("prereq"):
            para(doc, "Voraussetzungen / Warnhinweise", bold=True)
            add_bullets(doc, [fmt(p, ctx) for p in m["prereq"]])
        if m.get("nachweis"):
            para(doc, "Nachweis: " + fmt(m["nachweis"], ctx), italic=True)

    # False positives
    fps = [r for r in model["rows"] if r["status"] == FP]
    if fps:
        doc.add_heading(f"{sec}  Bekannte False Positives", level=1)
        sec += 1
        add_bullets(doc, [f"{r['cid']} {r['title']}: {r['befund']}" for r in fps])

    # Manual controls
    man = [r for r in model["rows"] if r["status"] == MANUAL]
    doc.add_heading(f"{sec}  Manuelle Kontrollen — CISO-Review ({len(man)})", level=1)
    sec += 1
    para(
        doc,
        "Prowler prüft diese Kontrollen nicht. Jede braucht eine Stichprobe nach dem Prüfweg der Checkliste (Spalte J); die Maßnahme nennt Kontext und Kommandos.",
    )

    def man_ref(r):
        if r["massnahme"] != "Review":
            return r["massnahme"]
        sib = [
            x
            for x in model["rows"]
            if x["cid"] in (by_m_all.get(r["cid"], {}).get("controls") or [])
            and x["pruefart"] == "Automated"
        ]
        note = (
            " — Ressourcentyp heute nicht vorhanden"
            if sib and all(x["status"] == NA for x in sib)
            else ""
        )
        return f"Review: {r.get('review_ref', '')}{note}"

    by_m_all = {c: m for m in CATALOGUE for c in m.get("controls", [])}
    add_table(
        doc,
        [["CIS #", "Kontrolle", "Level", "Maßnahme"]]
        + [[r["cid"], r["title"], r["level"], man_ref(r)] for r in man],
        widths=[1.6, 10.4, 2, 3],
        header=True,
    )

    # Wiedervorlage
    na = [r for r in model["rows"] if r["status"] == NA]
    if na:
        doc.add_heading(
            f"{sec}  Wiedervorlage — Kontrollen ohne Ressourcen ({len(na)})", level=1
        )
        sec += 1
        para(
            doc,
            "Der Ressourcentyp existiert heute nicht. Beim ersten Rollout (Cloud SQL, BigQuery, Dataproc, GKE, …) gilt die Kontrolle — das genannte Pack setzt die Org-Policy vorab, so dass der Rollout konform startet.",
        )
        by_m = {c: m for m in CATALOGUE for c in m.get("controls", [])}
        add_table(
            doc,
            [["CIS #", "Kontrolle", "Vorab durch Pack / Maßnahme"]]
            + [
                [
                    r["cid"],
                    r["title"],
                    ", ".join(by_m[r["cid"]].get("packs", []))
                    or by_m[r["cid"]]["title"]
                    if r["cid"] in by_m
                    else "—",
                ]
                for r in na
            ],
            widths=[1.6, 10.4, 5],
            header=True,
        )

    # Workloads — the projects beside the estate, each with its consumer check
    if "workloads" in scope:
        wl = model["workloads"]
        doc.add_heading(
            f"{sec}  Workloads — Projekte neben dem Estate ({len(wl)})", level=1
        )
        sec += 1
        if wl:
            para(
                doc,
                "Neben dem Estate stehen Projekte, die sein Interface lesen: ein satz-Projekt mit eigenem Estate, IaC-Service-Account und State, "
                "oder ein Terraform-Verzeichnis, das module.satz.<export> liest. Ihre Google-Projekte gehören zu den eigenen Projekten des Prowler-Scans. "
                "Der Consumer-Check hält jedes Projekt offline an die Regeln des Interfaces — Anhänge nur an Attach-Punkten, keine autoritative Rolle "
                "und keine Org-Policy auf einem Knoten des Estates, keine Ressource, die das Estate selbst deklariert; ein satz-Projekt wird zusätzlich "
                "kompiliert (transpile --check). Checkov läuft über das erzeugte HCL des Projekts.",
            )
            add_table(
                doc,
                [["Workload", "Art", "Google-Projekt", "Consumer-Check", "Checkov"]]
                + [
                    [
                        w["name"],
                        w["kind_de"],
                        w["project_id"] or "—",
                        w["check"],
                        w["checkov"],
                    ]
                    for w in wl
                ],
                widths=[2.4, 3.6, 3.4, 5.2, 2.4],
                header=True,
            )
            if any(w["check_ok"] is False for w in wl):
                para(
                    doc,
                    "Ein abgelehnter Consumer-Check ist eine Maßnahme des Projekt-Teams, nicht des Estates: das Projekt ändert seine Deklaration, "
                    "bis der Check ohne Befund durchläuft; das Estate bleibt unverändert.",
                    italic=True,
                )
        else:
            para(
                doc,
                "Neben dem Estate steht kein Workload: kein Projekt liest das Interface des Estates.",
            )

    # Anhang non-CIS
    extra = collections.Counter(
        f["finding_info"]["title"]
        for f in model["findings"]
        if f.get("status_code") == "FAIL" and not cis_ids(f) and proj_of(f) in set(own)
    )
    if extra:
        doc.add_heading(
            f"{sec}  Anhang — weitere Prowler-Befunde außerhalb CIS 5.0", level=1
        )
        sec += 1
        para(
            doc,
            "Best-Practice-Prüfungen ohne CIS-Zuordnung, nur eigene Projekte, zur Kenntnis — nicht Teil dieses Plans:",
        )
        add_bullets(doc, [f"{n} × {t}" for t, n in extra.most_common(15)])

    # Nachweis
    doc.add_heading(f"{sec}  Nachweis der Umsetzung", level=1)
    sec += 1
    para(
        doc,
        "Nach den Sofort- und kurzfristigen Maßnahmen wird derselbe Scan wiederholt — ein Prowler-Lauf je Service mit Deadline (ein monolithischer Org-Scan "
        "friert auf großen Organisationen ein), Compute auf den eigenen Projekten, dann satz report-compliance:",
    )
    add_code(
        doc,
        f"""SA={S.get("iac_sa", "<iac-sa>")}; ORG={ORG}; OUT=audit/{DOM}/$(date +%F)/raw
prowler_scan.sh --org $ORG --sa $SA --out $OUT --prefix {S["kunde"]} --own-projects "{" ".join(own)}"
merge_ocsf.py --raw $OUT --prefix {S["kunde"]}{(" --fleet-regex " + repr(S["fleet_regex"])) if S.get("fleet_regex") else ""}
satz_evidence.sh --repo {Path(S["estate_repo"]).name} --estate {S.get("estate_file", "")} --ocsf $OUT/{S["kunde"]}-prowler-merged.ocsf.json --out $OUT/../evidence
build_audit.py --scope scope.yaml --ocsf $OUT/{S["kunde"]}-prowler-merged.ocsf.json --satz-report $OUT/../evidence/satz-report-cis-gcp-5.0.json --out-dir audit""",
    )
    flips = collections.defaultdict(list)
    for m in measures:
        for r in m["_active"]:
            flips[m["num"]].append(r["cid"])
    if flips:
        para(
            doc,
            "Erwartetes Ergebnis: "
            + "; ".join(f"{k}: {', '.join(v)} → PASS" for k, v in flips.items())
            + ". Die MANUAL-Kontrollen bleiben Sache des CISO-Reviews; die N/A-Kontrollen sind beim ersten Rollout des Ressourcentyps erneut zu bewerten.",
        )
    doc.add_heading("Werkzeuggrenzen dieses Scans", level=2)
    add_bullets(
        doc,
        list(S.get("scan_caveats") or [])
        + [
            f"{sum(cnt.values())} Kontrollen, {n_own_f + n_fl_f} FAIL-Einzelbefunde — die Zahl der Befunde ist keine Zahl der Probleme; die Maßnahmen M1–M{len(measures)} bündeln sie.",
            "Prowler 5 emittiert keine MANUAL-Zeilen; die manuellen Kontrollen stammen aus dem CIS-Katalog. Prowlers CIS-5.0-Zuordnung führt 2.13 nicht (check_overrides in scope.yaml).",
        ],
    )
    doc.add_heading("Offene Punkte für das Review", level=2)
    ops = [
        fmt(m["open_point"], measure_ctx(scope, model, m))
        for m in measures
        if m.get("open_point")
    ] + list(S.get("open_points") or [])
    if fps:
        ops.append(
            "False Positives nach Umsetzung in der Audit-Akte dokumentieren: "
            + ", ".join(r["cid"] for r in fps)
            + "."
        )
    for w in model.get("workloads") or []:
        if w["check_ok"] is None:
            ops.append(
                f"Workload {w['name']}: kein Nachweis — satz_evidence.sh mit --workloads erneut laufen lassen und neu bauen."
            )
        elif w["check_ok"] is False:
            ops.append(
                f"Workload {w['name']}: Consumer-Check abgelehnt ({w['check']}) — mit dem Projekt-Team klären."
            )
    add_bullets(doc, ops or ["—"])
    doc.save(out)


CATALOGUE = []


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scope", required=True)
    ap.add_argument("--catalog")
    ap.add_argument("--measures")
    ap.add_argument("--ocsf", nargs="+", required=True)
    ap.add_argument("--satz-report")
    ap.add_argument("--satz-require")
    ap.add_argument(
        "--evidence-dir",
        help="where satz_evidence.sh wrote consumer-<workload>.txt and checkov-<workload>.json (default: the folder of --satz-report)",
    )
    ap.add_argument("--out-dir", required=True)
    a = ap.parse_args()
    global CATALOGUE
    scope = yaml.safe_load(open(Path(a.scope).expanduser()))
    for k in ("estate_repo", "audit_dir"):
        if scope.get(k):
            scope[k] = str(Path(scope[k]).expanduser())
    catalog = json.load(open(a.catalog or find_catalog()))
    CATALOGUE = yaml.safe_load(
        open(a.measures or HERE.parent / "assets" / "measures.yaml")
    )["measures"]
    findings = load_ocsf(a.ocsf)
    satz_rows = (
        json.load(open(Path(a.satz_report).expanduser())).get("rows", [])
        if a.satz_report
        else []
    )
    require = (
        json.load(open(Path(a.satz_require).expanduser())) if a.satz_require else None
    )
    model = build_model(scope, catalog, findings, satz_rows, require)
    evidence_dir = (
        Path(a.evidence_dir).expanduser()
        if a.evidence_dir
        else (Path(a.satz_report).expanduser().parent if a.satz_report else None)
    )
    model["workloads"] = load_workloads(scope, evidence_dir)
    measures = select_measures(scope, model, CATALOGUE)
    out = Path(a.out_dir).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    x = out / f"CIS-5.0-Checkliste_{scope['kunde']}_{scope['scan_date']}.xlsx"
    d = out / f"Remediation-Plan_{scope['kunde']}_CIS-5.0_{scope['scan_date']}.docx"
    build_xlsx(scope, model, measures, x)
    build_docx(scope, model, measures, d)
    print(
        f"findings: {len(findings)} (own projects {len(model['own'])}, fleet {len(model['fleet'])}; compute scanned: {model['compute_scanned']})"
    )
    print(
        "status counts:",
        {k: model["counts"][k] for k in STATUS_ORDER},
        "sum",
        sum(model["counts"].values()),
    )
    print(
        "measures:",
        "; ".join(
            f"{m['num']} {m['id']} [{m['phase']}] {','.join(r['cid'] for r in m['_active'])}"
            for m in measures
        ),
    )
    if "workloads" in scope:
        wl = model["workloads"]
        print(
            f"workloads: {len(wl)} (passed {sum(1 for w in wl if w['check_ok'])}, refused {sum(1 for w in wl if w['check_ok'] is False)}, "
            f"no evidence {sum(1 for w in wl if w['check_ok'] is None)})"
        )
    print("wrote", x)
    print("wrote", d)


if __name__ == "__main__":
    main()
